import json
import os
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from PIL import Image
import pytesseract
from PyPDF2 import PdfReader

from config import TESSERACT_CMD


def _configure_tesseract():
    """Config override -> env var -> PATH -> common Windows install."""
    candidates = [
        TESSERACT_CMD,
        os.environ.get("TESSERACT_CMD"),
        shutil.which("tesseract"),
        r"C:\Users\44969\AppData\Local\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    ]
    for c in candidates:
        if c and os.path.exists(c):
            pytesseract.pytesseract.tesseract_cmd = c
            return c
    return None


_configure_tesseract()


class FileReader:

    IMAGE_EXT = [".png", ".jpg", ".jpeg", ".tiff", ".bmp"]

    def read_file(self, file_path):

        file_path = Path(file_path)
        extension = file_path.suffix.lower()

        if extension == ".txt":
            return self._read_txt(file_path)
        elif extension == ".csv":
            return self._read_csv(file_path)
        elif extension == ".json":
            return self._read_json(file_path)
        elif extension == ".pdf":
            return self._read_pdf(file_path)
        elif extension == ".docx":
            return self._read_docx(file_path)
        elif extension in self.IMAGE_EXT:
            return self._read_image(file_path)
        else:
            raise ValueError(f"Unsupported file format: {extension}")

    @staticmethod
    def is_image_based(file_path):
        return Path(file_path).suffix.lower() in FileReader.IMAGE_EXT

    # ---------------- TXT
    def _read_txt(self, file_path):
        with open(file_path, "r", encoding="utf-8", errors="replace") as file:
            return file.read()

    # ---------------- CSV
    def _read_csv(self, file_path):
        df = pd.read_csv(file_path)
        return df.to_string(index=False)

    def read_csv_records(self, file_path):
        """Structured CSV read used by batch mode (keeps every column)."""
        return pd.read_csv(file_path)

    # ---------------- JSON
    def _read_json(self, file_path):
        with open(file_path, "r", encoding="utf-8") as file:
            data = json.load(file)
        return json.dumps(data, indent=2)

    # ---------------- PDF
    def _read_pdf(self, file_path):
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text

    # ---------------- DOCX
    def _read_docx(self, file_path):
        doc = Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs]
        for table in doc.tables:                      # tables can hold PHI too
            for row in table.rows:
                paragraphs.append(" ".join(c.text for c in row.cells))
        return "\n".join(paragraphs)

    # ---------------- IMAGE OCR
    def _read_image(self, file_path):
        image = Image.open(file_path).convert("L")
        return pytesseract.image_to_string(image, config="--psm 6")
